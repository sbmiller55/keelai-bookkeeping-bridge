"""
Revenue Recognition Schedule endpoints.

Routes:
  GET  /clients/{id}/revenue/summary
  GET  /clients/{id}/revenue/streams
  POST /clients/{id}/revenue/streams
  PATCH /clients/{id}/revenue/streams/{sid}
  DELETE /clients/{id}/revenue/streams/{sid}
  GET  /clients/{id}/revenue/contracts
  POST /clients/{id}/revenue/contracts
  PATCH /clients/{id}/revenue/contracts/{cid}
  DELETE /clients/{id}/revenue/contracts/{cid}
  POST /clients/{id}/revenue/contracts/{cid}/generate-jes
  GET  /clients/{id}/revenue/ar-aging
  GET  /clients/{id}/revenue/integration-settings
  PUT  /clients/{id}/revenue/integration-settings
  POST /clients/{id}/revenue/sync
"""

import calendar
import io
import json
from datetime import datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from auth import get_current_user
from database import get_db
import errors as errors_mod
from models import (
    BillingType,
    Client,
    JournalEntry,
    RevenueContract,
    RevenueContractStatus,
    RevenueIntegrationSettings,
    RevenueScheduleEntry,
    RevenueStream,
    Transaction,
    TransactionStatus,
    User,
    next_je_number,
)
from schemas import (
    RevenueContractCreate,
    RevenueContractRead,
    RevenueContractUpdate,
    RevenueIntegrationSettingsRead,
    RevenueIntegrationSettingsUpdate,
    RevenueScheduleEntryRead,
    RevenueSummary,
    RevenueStreamCreate,
    RevenueStreamRead,
    RevenueStreamUpdate,
)

router = APIRouter(prefix="/clients/{client_id}/revenue", tags=["revenue"])


# ── Helpers ───────────────────────────────────────────────────────────────────

def _get_client(client_id: int, user: User, db: Session) -> Client:
    c = db.query(Client).filter(Client.id == client_id, Client.user_id == user.id).first()
    if not c:
        raise HTTPException(404, "Client not found")
    return c


def _current_month() -> str:
    return datetime.utcnow().strftime("%Y-%m")


def _last_day(year: int, month: int) -> datetime:
    return datetime(year, month, calendar.monthrange(year, month)[1])


def _add_months(dt: datetime, n: int) -> datetime:
    month = dt.month + n
    year = dt.year + (month - 1) // 12
    month = (month - 1) % 12 + 1
    return dt.replace(year=year, month=month, day=1)


def _get_or_create_integration_settings(client_id: int, db: Session) -> RevenueIntegrationSettings:
    settings = db.query(RevenueIntegrationSettings).filter(
        RevenueIntegrationSettings.client_id == client_id
    ).first()
    if not settings:
        settings = RevenueIntegrationSettings(client_id=client_id)
        db.add(settings)
        db.flush()
    return settings


POLICY_FULL_MONTH      = "full_month"
POLICY_DAILY_PRORATION = "daily_proration"

POLICY_NOTE = {
    POLICY_FULL_MONTH:      "ASC 606 revenue recognition — full month policy applied.",
    POLICY_DAILY_PRORATION: "ASC 606 revenue recognition — daily proration policy applied.",
}


def _month_label(period: str) -> str:
    """"2026-10" -> "Oct 2026", for column headers."""
    return datetime.strptime(period, "%Y-%m").strftime("%b %Y")


def _subscription_mode(client_id: int, db: Session) -> bool:
    """Whether this client uses the subscription Revenue tab.

    Gates the contracts / recognition-grid / revenue-queue behaviour so that
    clients billing on the other models keep the schedule view unchanged.
    """
    settings = db.query(RevenueIntegrationSettings).filter(
        RevenueIntegrationSettings.client_id == client_id
    ).first()
    return bool(getattr(settings, "subscription_mode", False)) if settings else False


def _recognition_policy(client_id: int, db: Session) -> str:
    """The client's recognition policy, defaulting to full month."""
    settings = db.query(RevenueIntegrationSettings).filter(
        RevenueIntegrationSettings.client_id == client_id
    ).first()
    policy = getattr(settings, "revenue_recognition_policy", None) if settings else None
    return policy if policy in (POLICY_FULL_MONTH, POLICY_DAILY_PRORATION) else POLICY_FULL_MONTH


def _term_months(start: datetime, end: datetime) -> int:
    """Whole months in a service period.

    Counted as the smallest n where start + n months is past the end, rather
    than by counting the calendar months the period touches. The difference
    matters: a 12-month subscription running 4 Sep 2026 to 3 Sep 2027 touches
    thirteen calendar months, and counting those gives a 13-month schedule for
    a one-year subscription.
    """
    n = 1
    while _add_months_keep_day(start, n) <= end and n < 600:
        n += 1
    return n


def _add_months_keep_day(d: datetime, n: int) -> datetime:
    """Shift by whole months without snapping to the 1st (unlike _add_months)."""
    month = d.month + n
    year = d.year + (month - 1) // 12
    month = (month - 1) % 12 + 1
    day = min(d.day, calendar.monthrange(year, month)[1])
    return d.replace(year=year, month=month, day=day)


def _build_schedule(
    contract: RevenueContract,
    stream: Optional[RevenueStream],
    policy: str = POLICY_FULL_MONTH,
) -> list[dict]:
    """Generate the month-by-month recognition schedule entries for a contract.

    `policy` only decides what happens to a **single-month** subscription whose
    service period straddles two calendar months — a monthly plan starting on
    the 4th, say, which runs to the 3rd of the next month:

      full_month      100% in the calendar month the period starts. The
                      simplified treatment most SaaS companies use, acceptable
                      under ASC 606 when applied consistently.
      daily_proration split across both months by actual days of service.

    Multi-month contracts (annual plans) are unaffected by the policy: they
    already defer ratably across their term, and the policy question doesn't
    arise. They are still split evenly, one period per month of the term.
    """
    if not stream or not contract.service_period_start or not contract.service_period_end:
        return []

    sp_start = contract.service_period_start
    sp_end = contract.service_period_end
    value = round(contract.total_contract_value or 0.0, 2)
    n_months = _term_months(sp_start, sp_end)

    # ── Single-month subscription: the policy decides ────────────────────────
    if n_months == 1:
        if policy == POLICY_DAILY_PRORATION and sp_end.strftime("%Y-%m") != sp_start.strftime("%Y-%m"):
            total_days = (sp_end.date() - sp_start.date()).days + 1
            if total_days <= 0:
                return [{"period": sp_start.strftime("%Y-%m"), "amount": value}]
            first_month_end = _last_day(sp_start.year, sp_start.month)
            first_days = (first_month_end.date() - sp_start.date()).days + 1
            first_amt = round(value * first_days / total_days, 2)
            return [
                {"period": sp_start.strftime("%Y-%m"), "amount": first_amt},
                {"period": sp_end.strftime("%Y-%m"), "amount": round(value - first_amt, 2)},
            ]
        # Full month (and any single-month period that doesn't straddle).
        return [{"period": sp_start.strftime("%Y-%m"), "amount": value}]

    # ── Multi-month: ratable across the term, one period per month ───────────
    monthly = round(value / n_months, 2)
    periods = []
    current = sp_start.replace(day=1)
    total = 0.0
    for i in range(n_months):
        is_last = (i == n_months - 1)
        amt = round(value - total, 2) if is_last else monthly
        periods.append({"period": current.strftime("%Y-%m"), "amount": amt})
        total += amt
        current = _add_months(current, 1)
    return periods


def _enrich_contract(contract: RevenueContract, db: Session) -> dict:
    entries = db.query(RevenueScheduleEntry).filter(
        RevenueScheduleEntry.contract_id == contract.id
    ).order_by(RevenueScheduleEntry.period).all()
    d = {
        "id": contract.id,
        "client_id": contract.client_id,
        "revenue_stream_id": contract.revenue_stream_id,
        "customer_name": contract.customer_name,
        "external_id": contract.external_id,
        "source": contract.source,
        "invoice_number": contract.invoice_number,
        "total_contract_value": contract.total_contract_value,
        "billing_date": contract.billing_date.isoformat() if contract.billing_date else None,
        "due_date": contract.due_date.isoformat() if contract.due_date else None,
        "service_period_start": contract.service_period_start.isoformat() if contract.service_period_start else None,
        "service_period_end": contract.service_period_end.isoformat() if contract.service_period_end else None,
        "amount_recognized": contract.amount_recognized,
        "amount_deferred": round(contract.total_contract_value - contract.amount_recognized, 2),
        "payment_received": contract.payment_received,
        "payment_date": contract.payment_date.isoformat() if contract.payment_date else None,
        "status": contract.status.value if contract.status else "active",
        "ai_confidence": contract.ai_confidence,
        "ai_reasoning": contract.ai_reasoning,
        "created_at": contract.created_at.isoformat(),
        "schedule": [
            {"id": e.id, "period": e.period, "amount": e.amount, "je_id": e.je_id, "recognized": e.recognized}
            for e in entries
        ],
    }
    # The Contracts table reports MRR, which is the schedule's monthly rate —
    # an annual plan's $60 is $5 of MRR, not $60. Taken from the schedule the
    # contract actually has rather than recomputed from the term, so a contract
    # whose schedule was edited reports what it will really recognize.
    months = len(entries)
    d["term_months"] = months
    d["mrr"] = round(contract.total_contract_value / months, 2) if months else 0.0
    d["billing_cadence"] = "annual" if months > 1 else "monthly"
    return d


# ── Summary ───────────────────────────────────────────────────────────────────

@router.get("/summary")
def get_summary(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    this_month = _current_month()
    now = datetime.utcnow()

    contracts = db.query(RevenueContract).filter(
        RevenueContract.client_id == client_id
    ).all()

    # Revenue recognized in the month being reported on. Recognition is monthly
    # and posts at month end, so for most of any given month the current month
    # is legitimately zero — which read as "recognition is broken" on screen.
    # When that is the case, report the last month that does have recognition
    # and tell the page which month it is showing.
    def _recognized_for(period: str) -> float:
        return sum(
            e.amount for e in db.query(RevenueScheduleEntry).filter(
                RevenueScheduleEntry.client_id == client_id,
                RevenueScheduleEntry.period == period,
                RevenueScheduleEntry.recognized == True,
            ).all()
        )

    recognized_period = this_month
    recognized_this_month = _recognized_for(this_month)
    if not recognized_this_month:
        last_closed = (
            db.query(RevenueScheduleEntry.period)
            .filter(
                RevenueScheduleEntry.client_id == client_id,
                RevenueScheduleEntry.recognized == True,
                RevenueScheduleEntry.period < this_month,
            )
            .order_by(RevenueScheduleEntry.period.desc())
            .first()
        )
        if last_closed:
            recognized_period = last_closed[0]
            recognized_this_month = _recognized_for(recognized_period)

    # Total deferred: contracts where billing_type has deferred component, payment received but not fully recognized
    total_deferred = sum(
        max(0.0, c.total_contract_value - c.amount_recognized)
        for c in contracts
        if c.payment_received and c.status == RevenueContractStatus.active
    )

    # Total AR outstanding: contracts where payment not received
    total_ar = sum(
        c.total_contract_value - c.amount_recognized
        for c in contracts
        if not c.payment_received and c.status == RevenueContractStatus.active
    )

    # Overdue: AR contracts with due_date in the past
    overdue_count = sum(
        1 for c in contracts
        if not c.payment_received
        and c.due_date
        and c.due_date < now
        and c.status == RevenueContractStatus.active
    )

    return RevenueSummary(
        recognized_this_month=round(recognized_this_month, 2),
        recognized_period=recognized_period,
        total_deferred=round(total_deferred, 2),
        total_ar_outstanding=round(total_ar, 2),
        invoices_overdue=overdue_count,
    )


# ── Revenue Streams ───────────────────────────────────────────────────────────

@router.get("/streams")
def list_streams(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    streams = db.query(RevenueStream).filter(RevenueStream.client_id == client_id).all()
    return [RevenueStreamRead.model_validate(s) for s in streams]


@router.post("/streams")
def create_stream(
    client_id: int,
    body: RevenueStreamCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    try:
        BillingType(body.billing_type)
    except ValueError:
        raise HTTPException(400, f"Invalid billing_type: {body.billing_type}")
    stream = RevenueStream(client_id=client_id, **body.model_dump())
    db.add(stream)
    db.commit()
    db.refresh(stream)
    return RevenueStreamRead.model_validate(stream)


@router.patch("/streams/{stream_id}")
def update_stream(
    client_id: int,
    stream_id: int,
    body: RevenueStreamUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    stream = db.query(RevenueStream).filter(
        RevenueStream.id == stream_id, RevenueStream.client_id == client_id
    ).first()
    if not stream:
        raise HTTPException(404, "Revenue stream not found")
    for field, val in body.model_dump(exclude_none=True).items():
        setattr(stream, field, val)
    db.commit()
    db.refresh(stream)
    return RevenueStreamRead.model_validate(stream)


@router.delete("/streams/{stream_id}")
def delete_stream(
    client_id: int,
    stream_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    stream = db.query(RevenueStream).filter(
        RevenueStream.id == stream_id, RevenueStream.client_id == client_id
    ).first()
    if not stream:
        raise HTTPException(404, "Revenue stream not found")
    db.delete(stream)
    db.commit()
    return {"ok": True}


# ── Revenue Contracts ─────────────────────────────────────────────────────────

@router.get("/contracts")
def list_contracts(
    client_id: int,
    status: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    q = db.query(RevenueContract).filter(RevenueContract.client_id == client_id)
    if status:
        q = q.filter(RevenueContract.status == status)
    contracts = q.order_by(RevenueContract.billing_date.desc().nullslast()).all()
    return [_enrich_contract(c, db) for c in contracts]


@router.post("/contracts")
def create_contract(
    client_id: int,
    body: RevenueContractCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    contract = RevenueContract(client_id=client_id, **body.model_dump())
    db.add(contract)
    db.flush()

    # Auto-generate schedule entries if stream and service period are set
    if body.revenue_stream_id and body.service_period_start and body.service_period_end:
        stream = db.query(RevenueStream).filter(RevenueStream.id == body.revenue_stream_id).first()
        if stream:
            _create_schedule_entries(contract, stream, db)

    db.commit()
    db.refresh(contract)
    return _enrich_contract(contract, db)


@router.patch("/contracts/{contract_id}")
def update_contract(
    client_id: int,
    contract_id: int,
    body: RevenueContractUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    contract = db.query(RevenueContract).filter(
        RevenueContract.id == contract_id, RevenueContract.client_id == client_id
    ).first()
    if not contract:
        raise HTTPException(404, "Contract not found")
    changed = body.model_dump(exclude_none=True)
    for field, val in changed.items():
        if field == "status":
            contract.status = RevenueContractStatus(val)
        else:
            setattr(contract, field, val)
    # If the service period (or stream) changed, rebuild the recognition schedule
    # so the months match — but only when nothing has been generated yet, so we
    # never orphan already-generated journal entries.
    if (("service_period_start" in changed or "service_period_end" in changed
            or "revenue_stream_id" in changed)
            and contract.revenue_stream_id
            and contract.service_period_start and contract.service_period_end):
        stream = db.query(RevenueStream).filter(
            RevenueStream.id == contract.revenue_stream_id
        ).first()
        has_generated = (
            db.query(RevenueScheduleEntry)
            .filter(
                RevenueScheduleEntry.contract_id == contract.id,
                RevenueScheduleEntry.je_id.isnot(None),
            )
            .first()
        )
        if stream and not has_generated:
            _create_schedule_entries(contract, stream, db)
    db.commit()
    return _enrich_contract(contract, db)


@router.delete("/contracts/{contract_id}")
def delete_contract(
    client_id: int,
    contract_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    contract = db.query(RevenueContract).filter(
        RevenueContract.id == contract_id, RevenueContract.client_id == client_id
    ).first()
    if not contract:
        raise HTTPException(404, "Contract not found")
    db.query(RevenueScheduleEntry).filter(RevenueScheduleEntry.contract_id == contract_id).delete()
    # Cash-collection transactions point at this contract, and the foreign key
    # would block the delete. The transaction is still real money collected, so
    # it is unlinked rather than removed — it reappears in the revenue queue as
    # an unmatched payment, which is the honest state once the contract is gone.
    db.query(Transaction).filter(Transaction.revenue_contract_id == contract_id).update(
        {"revenue_contract_id": None}, synchronize_session=False
    )
    db.delete(contract)
    db.commit()
    return {"ok": True}


@router.post("/contracts/{contract_id}/generate-jes")
def generate_recognition_jes(
    client_id: int,
    contract_id: int,
    period: Optional[str] = None,  # "YYYY-MM" — if None, generate all pending
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Generate journal entries for pending recognition schedule entries.
    Creates JEs in the review queue (status=pending on the synthetic transaction).
    """
    client = _get_client(client_id, current_user, db)
    contract = db.query(RevenueContract).filter(
        RevenueContract.id == contract_id, RevenueContract.client_id == client_id
    ).first()
    if not contract:
        raise HTTPException(404, "Contract not found")
    stream = None
    if contract.revenue_stream_id:
        stream = db.query(RevenueStream).filter(RevenueStream.id == contract.revenue_stream_id).first()
    if not stream:
        raise HTTPException(400, "Contract has no revenue stream configured")

    entries_q = db.query(RevenueScheduleEntry).filter(
        RevenueScheduleEntry.contract_id == contract_id,
        RevenueScheduleEntry.je_id == None,
    )
    if period:
        entries_q = entries_q.filter(RevenueScheduleEntry.period == period)
    entries = entries_q.all()

    created = []
    for entry in entries:
        try:
            sp_dt = datetime.strptime(entry.period, "%Y-%m")
        except ValueError:
            continue

        billing_type = BillingType(stream.billing_type) if isinstance(stream.billing_type, str) else stream.billing_type

        if billing_type in (BillingType.monthly_arrears, BillingType.invoice_completion):
            # AR-based: DR AR / CR Revenue (recognition entry)
            je_date = _last_day(sp_dt.year, sp_dt.month)
            debit_acct = stream.ar_account
            credit_acct = stream.revenue_account
            memo = f"Revenue Recognition - {contract.customer_name} - {sp_dt.strftime('%b %Y')}"
        else:
            # Deferred-based: DR Deferred Revenue / CR Revenue
            je_date = _last_day(sp_dt.year, sp_dt.month)
            debit_acct = stream.deferred_revenue_account
            credit_acct = stream.revenue_account
            memo = f"Revenue Recognition - {contract.customer_name} - {sp_dt.strftime('%b %Y')}"

        # Create synthetic transaction
        tx = Transaction(
            client_id=client_id,
            date=je_date,
            description=memo,
            amount=entry.amount,
            status=TransactionStatus.pending,
            source="revenue",
        )
        db.add(tx)
        db.flush()

        je = JournalEntry(
            transaction_id=tx.id,
            je_number=next_je_number(db),
            debit_account=debit_acct,
            credit_account=credit_acct,
            amount=entry.amount,
            je_date=je_date,
            memo=memo[:80],
            customer_name=contract.customer_name,
            ai_confidence=1.0,
            ai_reasoning=(
                f"{POLICY_NOTE.get(_recognition_policy(client_id, db), POLICY_NOTE[POLICY_FULL_MONTH])} "
                f"{contract.customer_name}, {stream.name}, {entry.period}."
            ),
        )
        db.add(je)
        db.flush()

        entry.je_id = je.id
        # Both of these were left untouched here, so the revenue page's
        # "Recognized this month" card read $0.00 no matter how much had been
        # recognized, and the deferred balance only moved when the monthly
        # auto-release job happened to be the thing that created the entry.
        # Only a finished service month counts as earned.
        if je_date < datetime.utcnow():
            entry.recognized = True
            contract.amount_recognized = round(
                (contract.amount_recognized or 0.0) + entry.amount, 2
            )
        else:
            tx.status = TransactionStatus.scheduled
        created.append(entry.period)

    # Cash receipt for a paid contract (DR bank / CR AR).
    if _generate_payment_je(contract, stream, db):
        created.append("payment")
    _ensure_customer_on_jes(contract, db)

    db.commit()
    return {"created": created, "contract_id": contract_id}


class BulkMatchBody(BaseModel):
    stream_id: int


@router.post("/contracts/bulk-match")
def bulk_match_contracts(
    client_id: int,
    body: BulkMatchBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Assign all unmatched contracts to a stream and infer service periods."""
    _get_client(client_id, current_user, db)
    stream = db.query(RevenueStream).filter(
        RevenueStream.id == body.stream_id,
        RevenueStream.client_id == client_id,
    ).first()
    if not stream:
        raise HTTPException(404, "Revenue stream not found")

    billing_type = BillingType(stream.billing_type) if isinstance(stream.billing_type, str) else stream.billing_type

    contracts = db.query(RevenueContract).filter(
        RevenueContract.client_id == client_id,
        RevenueContract.revenue_stream_id == None,
        RevenueContract.status == RevenueContractStatus.active,
        RevenueContract.total_contract_value > 0,
    ).all()

    updated = 0
    for contract in contracts:
        contract.revenue_stream_id = stream.id

        # Infer service period from billing date for simple billing types
        if not contract.service_period_start and contract.billing_date:
            bd = contract.billing_date
            if billing_type in (BillingType.invoice_completion, BillingType.monthly_arrears):
                contract.service_period_start = bd.replace(day=1)
                contract.service_period_end = _last_day(bd.year, bd.month)

        if contract.service_period_start and contract.service_period_end:
            _create_schedule_entries(contract, stream, db)

        updated += 1

    db.commit()
    return {"updated": updated}


# How many contracts a deferred stream needs in one month before its recognition
# posts as a single pooled entry instead of one per customer.
GROUP_DEFERRED_ABOVE = 5


def _book_recognition_for_period(client_id: int, month: str, db: Session) -> dict:
    """Book revenue recognition for ONE service month across every matched contract.

    One month is the only shape anyone asks for at close — "book September" —
    which is why this replaced a "generate all" pass. That pass walked every
    unbooked entry on every contract, so on a subscription client a single press
    posted a full annual schedule: twelve months of revenue, eleven of them not
    yet earned.

    Deferred-revenue streams post as one pooled entry per stream for the month.
    The deferred liability is not customer-specific, and a subscription client
    has hundreds of $5 slices that would bury the review queue — the Chargebee
    path already grouped this way, and this brings the rest of the app in line.
    Below `GROUP_DEFERRED_ABOVE` contracts the entries stay per customer, so a
    client with a handful of contracts keeps the customer detail on the JE. AR
    streams are always per contract: there the customer *is* the point of the
    entry, because the balance it moves is that customer's receivable.

    A month that has not finished yet is not earned. Its entries are still
    created, dated month end, but parked as `scheduled` — out of the review
    queue until the daily job promotes them. That is what makes the "this month"
    button safe to press on the 12th.
    """
    period_dt = datetime.strptime(month, "%Y-%m")
    je_date = _last_day(period_dt.year, period_dt.month)
    is_earned = je_date < datetime.utcnow()
    policy_note = POLICY_NOTE.get(
        _recognition_policy(client_id, db), POLICY_NOTE[POLICY_FULL_MONTH]
    )

    rows = (
        db.query(RevenueScheduleEntry, RevenueContract, RevenueStream)
        .join(RevenueContract, RevenueContract.id == RevenueScheduleEntry.contract_id)
        .join(RevenueStream, RevenueStream.id == RevenueContract.revenue_stream_id)
        .filter(
            RevenueScheduleEntry.client_id == client_id,
            RevenueScheduleEntry.period == month,
            RevenueScheduleEntry.je_id == None,          # noqa: E711 — SQL NULL
            RevenueContract.status == RevenueContractStatus.active,
        )
        .all()
    )

    def _is_ar(stream: RevenueStream) -> bool:
        bt = BillingType(stream.billing_type) if isinstance(stream.billing_type, str) else stream.billing_type
        return bt in (BillingType.monthly_arrears, BillingType.invoice_completion)

    # Group the deferred streams that are large enough to pool.
    by_stream: dict[int, list] = {}
    for entry, contract, stream in rows:
        by_stream.setdefault(stream.id, []).append((entry, contract, stream))
    pooled_streams = {
        sid for sid, items in by_stream.items()
        if not _is_ar(items[0][2]) and len(items) > GROUP_DEFERRED_ABOVE
    }

    created = 0
    total_amount = 0.0

    def _post(debit: str, credit: str, amount: float, memo: str, description: str,
              customer: Optional[str], reasoning: str) -> JournalEntry:
        tx = Transaction(
            client_id=client_id,
            date=je_date,
            description=description,
            amount=amount,
            status=TransactionStatus.pending if is_earned else TransactionStatus.scheduled,
            source="revenue",
        )
        db.add(tx)
        db.flush()
        je = JournalEntry(
            transaction_id=tx.id,
            je_number=next_je_number(db),
            debit_account=debit,
            credit_account=credit,
            amount=amount,
            je_date=je_date,
            memo=memo[:80],
            customer_name=customer,
            ai_confidence=1.0,
            ai_reasoning=reasoning,
        )
        db.add(je)
        db.flush()
        return je

    def _mark(entry: RevenueScheduleEntry, contract: RevenueContract, je: JournalEntry):
        entry.je_id = je.id
        # Both of these drive the revenue page: `recognized` is what
        # "Recognized" sums, and `amount_recognized` is what the deferred
        # balance is netted against. Leaving either unset — which the
        # auto-release path used to do — showed $0.00 recognized and a deferred
        # balance that never moved, however many JEs had posted.
        if is_earned:
            entry.recognized = True
            contract.amount_recognized = round((contract.amount_recognized or 0.0) + entry.amount, 2)
            if contract.amount_recognized >= round(contract.total_contract_value or 0.0, 2) - 0.01:
                contract.status = RevenueContractStatus.fully_recognized

    for sid, items in by_stream.items():
        stream = items[0][2]
        label = period_dt.strftime("%b %Y")

        if sid in pooled_streams:
            amount = round(sum(e.amount for e, _, _ in items), 2)
            if amount <= 0:
                continue
            memo = f"Revenue recognition - {stream.name} - {label}"
            je = _post(
                stream.deferred_revenue_account, stream.revenue_account, amount, memo,
                f"{memo} ({len(items)} contracts)", None,
                f"{policy_note} {stream.name}, {month}: {len(items)} contracts earned a "
                f"month of service against cash collected up front. Posted as one entry "
                f"for the period.",
            )
            for entry, contract, _ in items:
                _mark(entry, contract, je)
            created += 1
            total_amount += amount
            continue

        for entry, contract, _ in items:
            if entry.amount <= 0:
                continue
            debit = stream.ar_account if _is_ar(stream) else stream.deferred_revenue_account
            memo = f"Revenue Recognition - {contract.customer_name} - {label}"
            je = _post(
                debit, stream.revenue_account, entry.amount, memo, memo,
                contract.customer_name,
                f"{policy_note} {contract.customer_name}, {stream.name}, {month}.",
            )
            _mark(entry, contract, je)
            created += 1
            total_amount += entry.amount

    return {
        "month": month,
        "created": created,
        "entries": len(rows),
        "amount": round(total_amount, 2),
        "earned": is_earned,
        "status": "pending" if is_earned else "scheduled",
    }


@router.post("/generate-jes")
def generate_period_jes(
    client_id: int,
    month: Optional[str] = None,     # "YYYY-MM" — defaults to the current month
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Book one month's revenue recognition into the review queue.

    `month` defaults to the current month. A month still in progress posts as
    `scheduled` and lands in the queue on its own at month end.
    """
    _get_client(client_id, current_user, db)
    target = month or _current_month()
    try:
        datetime.strptime(target, "%Y-%m")
    except ValueError:
        raise HTTPException(400, "month must be in YYYY-MM format")
    result = _book_recognition_for_period(client_id, target, db)
    # Cash receipts for paid AR contracts, which are not tied to a service month.
    payments = 0
    for contract in db.query(RevenueContract).filter(
        RevenueContract.client_id == client_id,
        RevenueContract.status == RevenueContractStatus.active,
        RevenueContract.revenue_stream_id.isnot(None),
    ).all():
        stream = db.query(RevenueStream).filter(
            RevenueStream.id == contract.revenue_stream_id
        ).first()
        if _generate_payment_je(contract, stream, db):
            payments += 1
        _ensure_customer_on_jes(contract, db)
    db.commit()
    result["payment_jes"] = payments
    result["errors"] = []
    return result


# ── Subscription mode: contracts, forward recognition grid, revenue queue ────
# These three endpoints back the subscription Revenue tab. They are readable by
# any client, but the UI only shows them when subscription_mode is on.

def _forward_periods(count: int, start: Optional[datetime] = None) -> list[str]:
    """`count` calendar months as "YYYY-MM", starting with the current month."""
    base = (start or datetime.utcnow()).replace(day=1)
    out = []
    for i in range(count):
        d = _add_months(base, i)
        out.append(d.strftime("%Y-%m"))
    return out


def _open_recognition_periods(client_id: int, db: Session) -> list[dict]:
    """Months that have begun and still have recognition left to book.

    Deliberately not limited to the grid's 12-month window: a month that was
    never generated stays bookable however far back it is, and that is exactly
    the month someone needs to catch up on.
    """
    this_month = _current_month()
    rows = (
        db.query(RevenueScheduleEntry.period, RevenueScheduleEntry.amount)
        .join(RevenueContract, RevenueContract.id == RevenueScheduleEntry.contract_id)
        .filter(
            RevenueScheduleEntry.client_id == client_id,
            RevenueScheduleEntry.je_id == None,            # noqa: E711 — SQL NULL
            RevenueScheduleEntry.period <= this_month,     # string compare is safe on YYYY-MM
            RevenueContract.status == RevenueContractStatus.active,
            RevenueContract.revenue_stream_id.isnot(None),
        )
        .all()
    )
    totals: dict[str, dict] = {}
    for period, amount in rows:
        agg = totals.setdefault(period, {"period": period, "entries": 0, "amount": 0.0})
        agg["entries"] += 1
        agg["amount"] = round(agg["amount"] + (amount or 0.0), 2)
    return sorted(totals.values(), key=lambda r: r["period"])


def _grid_data(client_id: int, months: int, db: Session) -> dict:
    periods = _forward_periods(months)
    contracts = (
        db.query(RevenueContract)
        .filter(
            RevenueContract.client_id == client_id,
            RevenueContract.status != RevenueContractStatus.cancelled,
            RevenueContract.revenue_stream_id.isnot(None),
        )
        .order_by(RevenueContract.customer_name)
        .all()
    )
    streams = {
        s.id: s for s in db.query(RevenueStream).filter(RevenueStream.client_id == client_id).all()
    }
    by_contract: dict[int, dict[str, float]] = {}
    for entry in (
        db.query(RevenueScheduleEntry)
        .filter(
            RevenueScheduleEntry.client_id == client_id,
            RevenueScheduleEntry.period.in_(periods),
        )
        .all()
    ):
        by_contract.setdefault(entry.contract_id, {})[entry.period] = entry.amount

    rows = []
    for c in contracts:
        cells = by_contract.get(c.id, {})
        # A contract with nothing in the window is not forward-looking revenue,
        # so it is left out rather than shown as a row of dashes.
        if not cells:
            continue
        stream = streams.get(c.revenue_stream_id)
        rows.append({
            "contract_id": c.id,
            "customer_name": c.customer_name,
            "plan": stream.name if stream else "",
            "service_period_end": c.service_period_end.isoformat() if c.service_period_end else None,
            # None, not 0.0 — the grid shows a dash for a month this contract
            # does not cover, which is different from a month worth nothing.
            "cells": [cells.get(p) for p in periods],
            "total": round(sum(cells.get(p) or 0.0 for p in periods), 2),
        })

    totals = [
        round(sum(r["cells"][i] or 0.0 for r in rows), 2) for i in range(len(periods))
    ]
    return {
        "periods": periods,
        "rows": rows,
        "totals": totals,
        "grand_total": round(sum(totals), 2),
        "open_periods": _open_recognition_periods(client_id, db),
    }


@router.get("/recognition-grid")
def recognition_grid(
    client_id: int,
    months: int = 12,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Forward-looking recognition: one row per contract, one column per month."""
    _get_client(client_id, current_user, db)
    months = max(1, min(months, 36))
    return _grid_data(client_id, months, db)


@router.get("/recognition-grid.xlsx")
def export_recognition_grid(
    client_id: int,
    months: int = 12,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """The same grid as an xlsx, matching the prepaid-schedule export's shape."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    _get_client(client_id, current_user, db)
    data = _grid_data(client_id, max(1, min(months, 36)), db)
    periods = data["periods"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue Recognition"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="334155")
    headers = ["Customer", "Plan"] + [_month_label(p) for p in periods] + ["Total"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    # Totals on row 2, directly under the headers, matching the on-screen grid.
    tl = ws.cell(row=2, column=1, value="Total")
    tl.font = Font(bold=True)
    for i, amt in enumerate(data["totals"], start=3):
        c = ws.cell(row=2, column=i, value=amt)
        c.number_format = "$#,##0.00"
        c.font = Font(bold=True)
    g = ws.cell(row=2, column=len(headers), value=data["grand_total"])
    g.number_format = "$#,##0.00"
    g.font = Font(bold=True)

    for row_idx, row in enumerate(data["rows"], start=3):
        ws.cell(row=row_idx, column=1, value=row["customer_name"])
        ws.cell(row=row_idx, column=2, value=row["plan"])
        for i, amt in enumerate(row["cells"], start=3):
            # Dash, not blank: a month outside the contract's term is a
            # deliberate "nothing here", and a reader cannot tell that from an
            # empty cell.
            cell = ws.cell(row=row_idx, column=i, value=round(amt, 2) if amt is not None else "—")
            if amt is not None:
                cell.number_format = "$#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="center")
        t = ws.cell(row=row_idx, column=len(headers), value=row["total"])
        t.number_format = "$#,##0.00"
        t.font = Font(bold=True)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    for col_idx in range(3, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 13
    ws.freeze_panes = "C3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"revenue-recognition-{datetime.utcnow().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# What kind of revenue entry a pending transaction holds. Chargebee's import
# brings in more than the two headline types — fees, payouts, refunds and
# guarantee releases all arrive as pending entries too, and they are surfaced
# here rather than dropped, because the revenue queue is now the only place a
# subscription client's revenue entries appear.
def _classify_revenue_je(tx: Transaction, jes: list, stream_deferred: set[str]) -> str:
    desc = (tx.description or "").lower()
    if tx.source == "revenue":
        return "recognition"
    # Checked before the refund test on purpose: the reserve-release entry is
    # described "Refund reserve released - N subscriptions", so a refund test
    # that ran first would claim it.
    if "guarantee lapsed" in desc or "reserve released" in desc:
        return "reserve_release"
    if "refund" in desc:
        return "refund"
    if "processing fee" in desc or "fees" in desc:
        return "fees"
    if "payout" in desc:
        return "payout"
    for je in jes:
        if je.credit_account in stream_deferred:
            return "cash_collection"
    return "other"


JE_TYPE_LABELS = {
    "cash_collection": "Cash collected",
    "recognition":     "Revenue recognition",
    "refund":          "Refund",
    "fees":            "Processor fees",
    "payout":          "Bank payout",
    "reserve_release": "Guarantee lapsed",
    "other":           "Other",
}


@router.get("/review-queue")
def revenue_review_queue(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Pending revenue entries for this client, for the Revenue tab's own queue.

    Cash collections come from Chargebee, recognition entries from the monthly
    Generate button. Scheduled (future-month) entries are excluded — they are
    not reviewable until the month they belong to has finished.
    """
    _get_client(client_id, current_user, db)
    txs = (
        db.query(Transaction)
        .filter(
            Transaction.client_id == client_id,
            Transaction.status == TransactionStatus.pending,
            Transaction.source.in_(("revenue", "chargebee")),
        )
        .order_by(Transaction.date.desc(), Transaction.id.desc())
        .all()
    )
    tx_ids = [t.id for t in txs]
    je_map: dict[int, list] = {}
    if tx_ids:
        for je in db.query(JournalEntry).filter(JournalEntry.transaction_id.in_(tx_ids)).all():
            je_map.setdefault(je.transaction_id, []).append(je)

    streams = db.query(RevenueStream).filter(RevenueStream.client_id == client_id).all()
    deferred_accounts = {s.deferred_revenue_account for s in streams if s.deferred_revenue_account}
    contract_ids = {t.revenue_contract_id for t in txs if t.revenue_contract_id}
    contracts = {
        c.id: c
        for c in db.query(RevenueContract).filter(RevenueContract.id.in_(contract_ids)).all()
    } if contract_ids else {}

    rows = []
    for tx in txs:
        jes = je_map.get(tx.id, [])
        je_type = _classify_revenue_je(tx, jes, deferred_accounts)
        contract = contracts.get(tx.revenue_contract_id) if tx.revenue_contract_id else None
        # Only a cash collection can be "unmatched" — a recognition entry is
        # built from a contract's own schedule, so it always has one.
        unmatched = je_type == "cash_collection" and contract is None
        rows.append({
            "transaction_id": tx.id,
            "date": tx.date.isoformat() if tx.date else None,
            "customer_name": (
                (jes[0].customer_name if jes and jes[0].customer_name else None)
                or tx.counterparty_name
                or (contract.customer_name if contract else None)
            ),
            "invoice_number": (
                tx.invoice_number
                or (contract.invoice_number or contract.external_id if contract else None)
            ),
            "period": tx.date.strftime("%Y-%m") if tx.date else None,
            "je_type": je_type,
            "je_type_label": JE_TYPE_LABELS.get(je_type, "Other"),
            "description": tx.description,
            "amount": round(tx.amount or 0.0, 2),
            "contract_id": contract.id if contract else None,
            "unmatched": unmatched,
            "warning": (
                "No matching contract found — please review before approving."
                if unmatched else None
            ),
            "journal_entries": [
                {
                    "id": je.id,
                    "je_number": je.je_number,
                    "debit_account": je.debit_account,
                    "credit_account": je.credit_account,
                    "amount": round(je.amount or 0.0, 2),
                    "memo": je.memo,
                    "customer_name": je.customer_name,
                }
                for je in jes
            ],
        })

    return {
        "count": len(rows),
        "unmatched_count": sum(1 for r in rows if r["unmatched"]),
        "rows": rows,
    }


class RevenueApproveBody(BaseModel):
    transaction_ids: Optional[list[int]] = None
    approve_all: bool = False


@router.post("/review-queue/approve")
def approve_revenue_queue(
    client_id: int,
    body: RevenueApproveBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Approve revenue entries so they move on to the Export tab."""
    client = _get_client(client_id, current_user, db)
    q = db.query(Transaction).filter(
        Transaction.client_id == client_id,
        Transaction.status == TransactionStatus.pending,
        Transaction.source.in_(("revenue", "chargebee")),
    )
    if not body.approve_all:
        ids = [int(i) for i in (body.transaction_ids or [])]
        if not ids:
            raise HTTPException(400, "Pass transaction_ids, or approve_all=true.")
        q = q.filter(Transaction.id.in_(ids))
    txs = q.all()

    # Same chart-of-accounts gate the main queue applies on approval, so a bad
    # account name is caught here instead of at QBO export.
    from routers.transactions import _coa_violations_for_transaction

    approved, blocked = 0, []
    now = datetime.utcnow()
    for tx in txs:
        bad = _coa_violations_for_transaction(tx, client, db)
        if bad:
            blocked.append({"transaction_id": tx.id, "bad_account_names": bad})
            continue
        tx.status = TransactionStatus.approved
        for je in db.query(JournalEntry).filter(JournalEntry.transaction_id == tx.id).all():
            je.approved_by = current_user.id
            je.approved_at = now
        approved += 1
    db.commit()
    return {"approved": approved, "blocked": blocked}


class BackfillLinksBody(BaseModel):
    dry_run: bool = True


@router.post("/backfill-contract-links")
def backfill_contract_links(
    client_id: int,
    body: BackfillLinksBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Link cash-collection transactions imported before the link existed.

    Matching is on customer name + billing date + amount rather than the
    invoice id, because the id was only ever stored on the contract: the
    transaction did not record it until this was added. Those three together
    identify a billing, with one wrinkle — the subscriber names repeat, so the
    same name can bill the same amount on the same day twice. Candidates are
    therefore consumed as they are used, which pairs duplicates off one-to-one
    instead of pointing both transactions at the first contract.

    Defaults to a dry run: it reports what it would link and changes nothing.
    """
    _get_client(client_id, current_user, db)

    unlinked = [
        t for t in db.query(Transaction).filter(
            Transaction.client_id == client_id,
            Transaction.source == "chargebee",
            Transaction.revenue_contract_id == None,          # noqa: E711 — SQL NULL
        ).order_by(Transaction.date, Transaction.id).all()
        # Only the billings are cash collections. Fees, payouts and the
        # guarantee-release entry are not a customer's payment and have no
        # contract to belong to.
        if "billing" in (t.description or "").lower()
    ]

    # Contracts not already claimed by a transaction, bucketed by the match key.
    claimed = {
        row[0] for row in db.query(Transaction.revenue_contract_id)
        .filter(
            Transaction.client_id == client_id,
            Transaction.revenue_contract_id.isnot(None),
        ).all()
    }
    buckets: dict[tuple, list[RevenueContract]] = {}
    for c in db.query(RevenueContract).filter(
        RevenueContract.client_id == client_id
    ).order_by(RevenueContract.id).all():
        if c.id in claimed or not c.billing_date:
            continue
        key = (
            (c.customer_name or "").strip().lower(),
            c.billing_date.date(),
            round(c.total_contract_value or 0.0, 2),
        )
        buckets.setdefault(key, []).append(c)

    linked, ambiguous, unmatched = 0, 0, 0
    examples: list[dict] = []
    for tx in unlinked:
        if not tx.date:
            unmatched += 1
            continue
        key = (
            (tx.counterparty_name or "").strip().lower(),
            tx.date.date(),
            round(abs(tx.amount or 0.0), 2),
        )
        candidates = buckets.get(key) or []
        if not candidates:
            unmatched += 1
            if len(examples) < 10:
                examples.append({
                    "transaction_id": tx.id, "result": "unmatched",
                    "customer": tx.counterparty_name,
                    "date": tx.date.strftime("%Y-%m-%d"),
                    "amount": round(abs(tx.amount or 0.0), 2),
                })
            continue
        if len(candidates) > 1:
            ambiguous += 1          # still linked, to the oldest unclaimed one
        contract = candidates.pop(0)
        if not body.dry_run:
            tx.revenue_contract_id = contract.id
            # Backfill the invoice id too, so the queue shows it and a future
            # pass can match on it directly.
            if not tx.invoice_number:
                tx.invoice_number = contract.invoice_number or contract.external_id
        linked += 1
        if len(examples) < 10:
            examples.append({
                "transaction_id": tx.id, "result": "linked",
                "contract_id": contract.id, "customer": tx.counterparty_name,
                "invoice_number": contract.invoice_number or contract.external_id,
            })

    if not body.dry_run:
        db.commit()
    return {
        "dry_run": body.dry_run,
        "candidates": len(unlinked),
        "linked": linked,
        "duplicate_name_matches": ambiguous,
        "unmatched": unmatched,
        "examples": examples,
    }


# ── AR Aging ──────────────────────────────────────────────────────────────────

@router.get("/ar-aging")
def get_ar_aging(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    now = datetime.utcnow()

    contracts = db.query(RevenueContract).filter(
        RevenueContract.client_id == client_id,
        RevenueContract.payment_received == False,
        RevenueContract.status == RevenueContractStatus.active,
    ).order_by(RevenueContract.due_date.asc().nullslast()).all()

    result = []
    for c in contracts:
        days_out = (now - c.billing_date).days if c.billing_date else 0
        if c.due_date:
            days_past_due = (now - c.due_date).days
        else:
            days_past_due = days_out

        if days_past_due <= 0:
            aging_bucket = "current"
        elif days_past_due <= 30:
            aging_bucket = "1-30"
        elif days_past_due <= 60:
            aging_bucket = "31-60"
        elif days_past_due <= 90:
            aging_bucket = "61-90"
        else:
            aging_bucket = "over-90"

        result.append({
            "id": c.id,
            "customer_name": c.customer_name,
            "invoice_number": c.invoice_number,
            "billing_date": c.billing_date.isoformat() if c.billing_date else None,
            "due_date": c.due_date.isoformat() if c.due_date else None,
            "amount": c.total_contract_value,
            "days_outstanding": days_out,
            "days_past_due": max(0, days_past_due),
            "aging_bucket": aging_bucket,
            "source": c.source,
        })

    return result


# ── Integration Settings ──────────────────────────────────────────────────────

@router.get("/integration-settings")
def get_integration_settings(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    settings = _get_or_create_integration_settings(client_id, db)
    db.commit()
    return {
        "client_id": settings.client_id,
        "mercury_revenue_enabled": settings.mercury_revenue_enabled,
        "stripe_enabled": settings.stripe_enabled,
        "stripe_api_key": "***" if settings.stripe_api_key else None,
        "billcom_enabled": settings.billcom_enabled,
        "billcom_username": settings.billcom_username,
        "billcom_org_id": settings.billcom_org_id,
        "billcom_dev_key": "***" if settings.billcom_dev_key else None,
        "last_stripe_sync": settings.last_stripe_sync.isoformat() if settings.last_stripe_sync else None,
        "last_billcom_sync": settings.last_billcom_sync.isoformat() if settings.last_billcom_sync else None,
        # Reported as the effective value, not the raw column: rows created
        # before this setting existed hold NULL, and the engine treats NULL as
        # full month, so the UI must show the same thing it will actually do.
        "revenue_recognition_policy": _recognition_policy(client_id, db),
        # Same NULL-is-the-default reasoning as the policy above: rows created
        # before this column existed hold NULL, which means off.
        "subscription_mode": _subscription_mode(client_id, db),
    }


@router.put("/integration-settings")
def update_integration_settings(
    client_id: int,
    body: RevenueIntegrationSettingsUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_client(client_id, current_user, db)
    settings = _get_or_create_integration_settings(client_id, db)
    payload = body.model_dump(exclude_none=True)
    policy = payload.get("revenue_recognition_policy")
    if policy is not None and policy not in (POLICY_FULL_MONTH, POLICY_DAILY_PRORATION):
        raise HTTPException(
            422,
            f"revenue_recognition_policy must be one of "
            f"'{POLICY_FULL_MONTH}' or '{POLICY_DAILY_PRORATION}'.",
        )
    for field, val in payload.items():
        setattr(settings, field, val)
    db.commit()
    return {"ok": True}


# ── Sync ──────────────────────────────────────────────────────────────────────

@router.post("/sync")
def sync_revenue_sources(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Pull data from all enabled revenue sources (Mercury, Stripe, Bill.com),
    create RevenueContract records for new invoices, and run AI analysis to
    match them to revenue streams and set service periods.
    """
    client = _get_client(client_id, current_user, db)
    settings = _get_or_create_integration_settings(client_id, db)
    db.commit()

    streams = db.query(RevenueStream).filter(
        RevenueStream.client_id == client_id, RevenueStream.active == True
    ).all()
    streams_dicts = [
        {"id": s.id, "name": s.name, "billing_type": s.billing_type.value if hasattr(s.billing_type, "value") else s.billing_type,
         "revenue_account": s.revenue_account, "deferred_revenue_account": s.deferred_revenue_account}
        for s in streams
    ]

    imported = []
    errors = []

    # ── Mercury ───────────────────────────────────────────────────────────────
    if settings.mercury_revenue_enabled:
        try:
            imported += _sync_mercury_revenue(client_id, client, streams_dicts, db)
        except Exception as exc:
            errors.append(f"Mercury: {exc}")

    # ── Stripe ────────────────────────────────────────────────────────────────
    if settings.stripe_enabled and settings.stripe_api_key:
        try:
            imported += _sync_stripe(client_id, settings.stripe_api_key, client, streams_dicts, db)
            settings.last_stripe_sync = datetime.utcnow()
        except Exception as exc:
            errors.append(f"Stripe: {exc}")

    # ── Bill.com ──────────────────────────────────────────────────────────────
    if settings.billcom_enabled and settings.billcom_username and settings.billcom_password:
        try:
            imported += _sync_billcom(
                client_id, settings.billcom_username, settings.billcom_password,
                settings.billcom_org_id or "", settings.billcom_dev_key or "",
                client, streams_dicts, db,
            )
            settings.last_billcom_sync = datetime.utcnow()
        except Exception as exc:
            errors.append(f"Bill.com: {exc}")

    # ── Chargebee (subscription billings) ────────────────────────────────────
    # Unlike the sources above, this one returns its own summary: a subscription
    # import produces billing, reserve, refund, fee and payout entries, and the
    # bare contract count wouldn't describe what landed in the review queue.
    chargebee_result = None
    if getattr(settings, "chargebee_enabled", False):
        try:
            import chargebee_sync
            chargebee_result = chargebee_sync.run_import(
                client_id, db, getattr(settings, "chargebee_cursor", None)
            )
            if not chargebee_result.get("exhausted"):
                settings.chargebee_cursor = chargebee_result["cursor"]
        except Exception as exc:
            errors.append(f"Chargebee: {errors_mod.safe_detail(exc, 'import Chargebee activity')}")

    db.commit()
    return {
        "imported": len(imported),
        "contracts": imported,
        "chargebee": chargebee_result,
        "errors": errors,
    }


class ChargebeeRewindBody(BaseModel):
    days: int = 30


@router.post("/chargebee/setup")
def chargebee_setup(
    client_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Prepare a client for the Chargebee subscription demo.

    Creates the four plan streams and backfills a stretch of already-posted
    history so deferred-revenue balances and a part-recognized schedule are on
    screen from the start. History is marked exported, leaving the review queue
    empty for the live import to fill. Safe to re-run.
    """
    _get_client(client_id, current_user, db)
    import chargebee_sync
    settings = _get_or_create_integration_settings(client_id, db)
    result = chargebee_sync.setup_demo(client_id, db)
    settings.chargebee_enabled = True
    settings.chargebee_cursor = result["cursor"]
    db.commit()
    return result


@router.post("/chargebee/rewind")
def chargebee_rewind(
    client_id: int,
    body: ChargebeeRewindBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Move the demo cursor back so the same activity can be imported again."""
    _get_client(client_id, current_user, db)
    import chargebee_sync
    settings = _get_or_create_integration_settings(client_id, db)
    settings.chargebee_cursor = chargebee_sync.rewind_cursor(
        body.days, getattr(settings, "chargebee_cursor", None)
    )
    db.commit()
    return {"cursor": settings.chargebee_cursor, "rewound_days": body.days}


# ── Private sync helpers ──────────────────────────────────────────────────────

def _upsert_contract(
    client_id: int, raw_contract: dict, client, streams_dicts: list, db: Session
) -> Optional[dict]:
    """Create or update a RevenueContract from raw import data. Returns the external_id if new."""
    external_id = raw_contract.get("external_id")
    if external_id:
        existing = db.query(RevenueContract).filter(
            RevenueContract.client_id == client_id,
            RevenueContract.external_id == external_id,
        ).first()
        if existing:
            # Update payment status
            if raw_contract.get("payment_received") and not existing.payment_received:
                existing.payment_received = True
            # Backfill the payment date whenever we now have one and it's missing
            # (e.g. a contract already marked paid before we fetched dates).
            if raw_contract.get("payment_date") and not existing.payment_date:
                existing.payment_date = raw_contract.get("payment_date")
            return None  # already exists

    # AI analysis
    import ai_coder
    analysis = ai_coder.analyze_revenue_contract(raw_contract, streams_dicts, client)

    service_start = None
    service_end = None
    if analysis.get("service_period_start"):
        try:
            service_start = datetime.strptime(analysis["service_period_start"][:10], "%Y-%m-%d")
        except Exception:
            pass
    if analysis.get("service_period_end"):
        try:
            service_end = datetime.strptime(analysis["service_period_end"][:10], "%Y-%m-%d")
        except Exception:
            pass

    # Fallback: infer from billing_date if AI didn't set it
    if not service_start and raw_contract.get("period_start"):
        service_start = raw_contract["period_start"]
    if not service_end and raw_contract.get("period_end"):
        service_end = raw_contract["period_end"]

    contract = RevenueContract(
        client_id=client_id,
        revenue_stream_id=analysis.get("revenue_stream_id"),
        customer_name=raw_contract.get("customer_name", "Unknown"),
        external_id=external_id,
        source=raw_contract.get("source", "manual"),
        invoice_number=raw_contract.get("invoice_number"),
        total_contract_value=raw_contract.get("total_contract_value", 0.0),
        billing_date=raw_contract.get("billing_date"),
        due_date=raw_contract.get("due_date"),
        service_period_start=service_start,
        service_period_end=service_end,
        payment_received=raw_contract.get("payment_received", False),
        payment_date=raw_contract.get("payment_date"),
        ai_confidence=analysis.get("confidence"),
        ai_reasoning=analysis.get("reasoning"),
        raw_data=json.dumps(raw_contract.get("raw") or {}, default=str)[:4000],
        status=RevenueContractStatus.active,
    )
    db.add(contract)
    db.flush()

    # Generate schedule entries if we have all needed data
    if contract.revenue_stream_id and service_start and service_end:
        stream = next((s for s in db.query(RevenueStream).filter(RevenueStream.id == contract.revenue_stream_id).all()), None)
        if stream:
            _create_schedule_entries(contract, stream, db)

    return external_id


def _create_schedule_entries(contract: RevenueContract, stream: RevenueStream, db: Session):
    """Generate RevenueScheduleEntry rows for a contract."""
    # Remove any existing entries
    db.query(RevenueScheduleEntry).filter(
        RevenueScheduleEntry.contract_id == contract.id
    ).delete()

    periods = _build_schedule(contract, stream, _recognition_policy(contract.client_id, db))
    for p in periods:
        entry = RevenueScheduleEntry(
            contract_id=contract.id,
            client_id=contract.client_id,
            period=p["period"],
            amount=p["amount"],
            recognized=False,
        )
        db.add(entry)


def _generate_payment_je(contract: RevenueContract, stream: Optional[RevenueStream], db: Session) -> bool:
    """Book the AR cash receipt for a PAID revenue contract: DR bank / CR AR,
    dated on the payment date (falling back to due/billing date). Idempotent via
    contract.payment_je_id; requires the stream's bank_account to be set; and only
    fires once revenue has actually been recognized (AR debited) so we never push
    AR negative. Returns True if a receipt JE was created."""
    if not contract.payment_received or contract.payment_je_id:
        return False
    if not stream or not stream.bank_account:
        return False
    amount = abs(contract.total_contract_value or 0)
    if amount <= 0:
        return False
    has_recognition = (
        db.query(RevenueScheduleEntry)
        .filter(
            RevenueScheduleEntry.contract_id == contract.id,
            RevenueScheduleEntry.je_id.isnot(None),
        )
        .first()
    )
    if not has_recognition:
        return False
    pay_date = contract.payment_date or contract.due_date or contract.billing_date or datetime.utcnow()
    tx = Transaction(
        client_id=contract.client_id,
        date=pay_date,
        description=f"Payment received - {contract.customer_name} - {contract.invoice_number or ''}"[:255],
        amount=amount,
        status=TransactionStatus.pending,
        source="revenue",
    )
    db.add(tx)
    db.flush()
    je = JournalEntry(
        transaction_id=tx.id,
        je_number=next_je_number(db),
        debit_account=stream.bank_account,
        credit_account=stream.ar_account,
        amount=amount,
        je_date=pay_date,
        memo=f"Cash receipt: {contract.customer_name} inv {contract.invoice_number or ''}"[:80],
        customer_name=contract.customer_name,
        ai_confidence=1.0,
        ai_reasoning=f"AR cash receipt for {contract.customer_name} (paid per {contract.source}).",
    )
    db.add(je)
    db.flush()
    contract.payment_je_id = je.id
    return True


def _ensure_customer_on_jes(contract: RevenueContract, db: Session) -> int:
    """Set customer_name = contract.customer_name on every JE tied to this contract
    (recognition entries + cash receipt). QBO requires a customer on AR/income
    lines. Backfills existing JEs and is idempotent. Returns the count updated."""
    if not contract.customer_name:
        return 0
    je_ids = {
        e.je_id for e in db.query(RevenueScheduleEntry)
        .filter(
            RevenueScheduleEntry.contract_id == contract.id,
            RevenueScheduleEntry.je_id.isnot(None),
        ).all()
    }
    if contract.payment_je_id:
        je_ids.add(contract.payment_je_id)
    if not je_ids:
        return 0
    n = 0
    for je in db.query(JournalEntry).filter(JournalEntry.id.in_(list(je_ids))).all():
        if je.customer_name != contract.customer_name:
            je.customer_name = contract.customer_name
            n += 1
    return n


def _sync_mercury_revenue(client_id: int, client, streams_dicts: list, db: Session) -> list:
    """Pull Mercury transactions that look like revenue (incoming payments)."""
    # Incoming positive transactions sourced from Mercury that aren't internal transfers
    cutoff = datetime.utcnow() - timedelta(days=365)
    txns = db.query(Transaction).filter(
        Transaction.client_id == client_id,
        Transaction.amount > 0,
        Transaction.source == "mercury",
        Transaction.date >= cutoff,
        Transaction.status != TransactionStatus.transfer,
    ).all()

    imported = []
    for tx in txns:
        # Skip if already linked to a contract
        existing = db.query(RevenueContract).filter(
            RevenueContract.client_id == client_id,
            RevenueContract.external_id == tx.mercury_transaction_id,
        ).first()
        if existing:
            continue

        raw = {
            "external_id": tx.mercury_transaction_id or f"mercury-{tx.id}",
            "source": "mercury",
            "customer_name": tx.counterparty_name or tx.description or "Unknown",
            "invoice_number": tx.invoice_number,
            "total_contract_value": tx.amount,
            "billing_date": tx.date,
            "payment_received": True,
            "payment_date": tx.date,
            "description": tx.description or "",
            "raw": {"invoice_text": (tx.invoice_text or "")[:200]},
        }
        ext_id = _upsert_contract(client_id, raw, client, streams_dicts, db)
        if ext_id:
            imported.append(ext_id)
    return imported


def _sync_stripe(client_id: int, api_key: str, client, streams_dicts: list, db: Session) -> list:
    import stripe_client
    invoices = stripe_client.get_invoices(api_key)
    imported = []
    for inv in invoices:
        ext_id = _upsert_contract(client_id, inv, client, streams_dicts, db)
        if ext_id:
            imported.append(ext_id)
    return imported


def _sync_billcom(client_id: int, username: str, password: str, org_id: str, dev_key: str,
                  client, streams_dicts: list, db: Session) -> list:
    import billcom_client
    invoices = billcom_client.get_billcom_invoices(username, password, org_id, dev_key)
    imported = []
    for inv in invoices:
        ext_id = _upsert_contract(client_id, inv, client, streams_dicts, db)
        if ext_id:
            imported.append(ext_id)
    return imported
