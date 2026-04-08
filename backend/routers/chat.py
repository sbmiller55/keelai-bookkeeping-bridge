"""Claude AI chat assistant with rich client context."""
import json
import os
from pathlib import Path
from typing import Optional

import anthropic
from fastapi import APIRouter, Depends, HTTPException, status
from pydantic import BaseModel
from sqlalchemy.orm import Session

from auth import get_current_user
from database import get_db
import models
import storage

router = APIRouter(prefix="/chat", tags=["chat"])
_VOWELS = set("aeiou")
_COMMON_WORDS = {"in", "is", "of", "to", "at", "on", "an", "as", "or", "and", "the", "by", "for"}


class ChatMessage(BaseModel):
    role: str
    content: str


class ChatImage(BaseModel):
    data: str        # base64-encoded bytes
    media_type: str  # e.g. "image/png", "image/jpeg"


class ChatRequest(BaseModel):
    client_id: int
    messages: list[ChatMessage]
    current_page: Optional[str] = None   # e.g. "review_queue", "transactions"
    page_context: Optional[str] = None   # JSON string of data currently on screen
    images: list[ChatImage] = []         # attached to the latest user message


class ChatResponse(BaseModel):
    reply: str


# ── File readers ──────────────────────────────────────────────────────────────

def _read_coa_pdf(p: Path) -> Optional[str]:
    """Extract clean account list from a QBO Chart of Accounts PDF."""
    try:
        import pdfplumber
        accounts: list[str] = []
        with pdfplumber.open(p) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables():
                    for row in table:
                        if not row or len(row) < 2 or not row[1]:
                            continue
                        parts = row[1].split("\n")
                        joined = parts[0]
                        for part in parts[1:]:
                            s = part.strip()
                            if not s:
                                continue
                            ft = s.split()[0] if s.split() else ""
                            pl = joined[-1] if joined else ""
                            is_suffix = (
                                ft and len(ft) <= 3 and ft.islower()
                                and ft not in _COMMON_WORDS
                                and pl.islower() and pl not in _VOWELS
                            )
                            joined += s if is_suffix else " " + s
                        name = " ".join(joined.split()).strip()
                        if name.lower() not in ("name", "account", ""):
                            accounts.append(name)
        return "\n".join(sorted(set(accounts))) if accounts else None
    except Exception:
        return None


def _read_docx(p: Path) -> Optional[str]:
    try:
        import docx
        doc = docx.Document(p)
        text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        return text[:12000] if text else None
    except Exception:
        return None


def _read_file(path: Optional[str]) -> Optional[str]:
    if not path:
        return None
    with storage.as_local_path(path) as p:
        if p is None:
            return None
        suffix = p.suffix.lower()
        if suffix == ".pdf":
            return _read_coa_pdf(p)
        elif suffix in (".docx", ".doc"):
            return _read_docx(p)
        try:
            return p.read_text(errors="replace")[:12000]
        except Exception:
            return None


# ── System prompt builder ─────────────────────────────────────────────────────

def _build_system_prompt(
    client: models.Client,
    db: Session,
    current_page: Optional[str],
    page_context: Optional[str],
) -> str:
    parts = [
        f"You are an expert bookkeeping assistant for {client.name}. "
        "You have full access to their chart of accounts, accounting policy, all transactions, "
        "journal entries, rules engine, and any data currently visible on the page the user is viewing. "
        "Be concise and specific. Use accounting terminology correctly. "
        "When referencing amounts use dollar signs and two decimal places. "
        "When you see a correction (a human changed the AI's coding), note what was changed and why it matters.\n\n"
        "## Critical behavior rules\n"
        "1. NEVER claim or imply that you have done something unless you have already completed it via a tool call "
        "in this exact response. Do not say 'I've updated the COA' or 'I created the rule' unless the tool call "
        "already succeeded. If you intend to do something, do it first with the tool, then confirm. "
        "If a tool call fails, say so honestly.\n"
        "2. NEVER ask the user for a journal entry ID, transaction ID, or rule ID. "
        "You have all of this data in your context under 'All Transactions with Journal Entries'. "
        "Look it up yourself using the description, date, amount, or account names the user provides.\n\n"
        "## Rules Engine Access\n"
        "You can CREATE, LIST, and DISABLE rules that automatically code future transactions. "
        "When a user asks you to set up a rule — including excluding/rejecting certain transactions — use your tools to do it immediately. "
        "Do NOT tell the user to set up rules manually; create them yourself.\n\n"
        "Rule match types: description_contains, description_exact, counterparty_contains, counterparty_exact, "
        "category_equals, has_category, kind, amount_gt, amount_lt.\n\n"
        "Rule actions:\n"
        "  - expense: standard DR/CR journal entry\n"
        "  - prepaid: amortize over service period\n"
        "  - fixed_asset: capitalize and depreciate\n"
        "  - reject: mark transaction as rejected — NO journal entry, excluded from QBO export. "
        "Use this when the user wants to exclude/skip certain transactions (e.g. payroll handled elsewhere).\n\n"
        "Placeholders: $source_account (the Mercury account the txn came from), $category (Mercury category).\n\n"
        "After creating each rule, call apply_rule to apply it immediately to matching pending transactions.\n\n"
        "## Chart of Accounts\n"
        "You can add accounts to the Chart of Accounts using the update_chart_of_accounts tool. "
        "Use this when the user asks you to add accounts, or when a transaction needs an account that doesn't exist yet. "
        "Never tell the user to manually update the COA — do it yourself with the tool.",
    ]

    # Chart of accounts
    coa = _read_file(client.chart_of_accounts_path)
    if coa:
        parts.append(f"\n\n## Chart of Accounts\n{coa}")

    # Accounting policy
    policy = _read_file(client.policy_path)
    if policy:
        parts.append(f"\n\n## Accounting Policy\n{policy}")

    # All transactions with journal entries
    txns = (
        db.query(models.Transaction)
        .filter(models.Transaction.client_id == client.id)
        .order_by(models.Transaction.date.desc())
        .all()
    )
    jes_by_txn: dict[int, list[models.JournalEntry]] = {}
    all_jes = (
        db.query(models.JournalEntry)
        .join(models.Transaction, models.JournalEntry.transaction_id == models.Transaction.id)
        .filter(models.Transaction.client_id == client.id)
        .all()
    )
    for je in all_jes:
        jes_by_txn.setdefault(je.transaction_id, []).append(je)

    if txns:
        rows = []
        for t in txns:
            je_list = jes_by_txn.get(t.id, [])
            je_strs = []
            for je in je_list:
                conf = f"{int(je.ai_confidence * 100)}%" if je.ai_confidence is not None else "?"
                je_strs.append(
                    f"  je_id={je.id}"
                    + (f" je_number={je.je_number}" if je.je_number else "")
                    + f" DR:{je.debit_account} / CR:{je.credit_account} "
                    f"${je.amount:,.2f} conf={conf}"
                    + (f" memo={je.memo!r}" if je.memo else "")
                    + (f" je_date={je.je_date.strftime('%Y-%m-%d')}" if je.je_date else "")
                    + (f" reasoning={je.ai_reasoning[:80]!r}" if je.ai_reasoning else "")
                )
            je_block = "\n".join(je_strs) if je_strs else "  [no journal entry]"
            rows.append(
                f"txn_id={t.id} | {t.date.strftime('%Y-%m-%d')} | {t.description[:50]} | "
                f"${t.amount:,.2f} | {t.status} | {t.mercury_account_name or ''}\n{je_block}"
            )
        parts.append("\n\n## All Transactions with Journal Entries (newest first)\n" + "\n---\n".join(rows))

    # Approved transactions = reviewed/accepted by user (correction history)
    approved = [t for t in txns if t.status == "approved"]
    if approved:
        summary_lines = []
        for t in approved[:30]:
            je_list = jes_by_txn.get(t.id, [])
            for je in je_list:
                summary_lines.append(
                    f"{t.description[:40]} → DR:{je.debit_account} / CR:{je.credit_account}"
                )
        if summary_lines:
            parts.append(
                "\n\n## Human-Approved Codings (these represent confirmed correct account choices)\n"
                + "\n".join(summary_lines)
            )

    # Current page context from frontend
    if current_page and page_context:
        page_labels = {
            "review_queue": "Review Queue (pending transactions currently on screen)",
            "transactions": "Transactions page",
            "rules": "Rules page",
            "export": "Export page",
        }
        label = page_labels.get(current_page, current_page)
        try:
            parsed = json.loads(page_context)
            formatted = json.dumps(parsed, indent=2)
        except Exception:
            formatted = page_context
        parts.append(f"\n\n## Current Page: {label}\nThis is the exact data the user is looking at right now:\n{formatted}")

    return "\n".join(parts)


# ── Tools Claude can call ─────────────────────────────────────────────────────

_TOOLS = [
    {
        "name": "create_journal_entry",
        "description": "Create a new journal entry for an existing transaction. Use this to add a split or an entirely new JE.",
        "input_schema": {
            "type": "object",
            "properties": {
                "transaction_id": {"type": "integer", "description": "The transaction ID to attach this JE to"},
                "debit_account": {"type": "string"},
                "credit_account": {"type": "string"},
                "amount": {"type": "number"},
                "memo": {"type": "string"},
                "je_date": {"type": "string", "description": "ISO date string YYYY-MM-DD (optional)"},
            },
            "required": ["transaction_id", "debit_account", "credit_account", "amount"],
        },
    },
    {
        "name": "update_journal_entry",
        "description": "Update a journal entry's debit account, credit account, amount, memo, or je_date.",
        "input_schema": {
            "type": "object",
            "properties": {
                "je_id": {"type": "integer", "description": "The journal entry ID"},
                "debit_account": {"type": "string"},
                "credit_account": {"type": "string"},
                "amount": {"type": "number"},
                "memo": {"type": "string"},
                "je_date": {"type": "string", "description": "ISO date string YYYY-MM-DD"},
            },
            "required": ["je_id"],
        },
    },
    {
        "name": "delete_journal_entry",
        "description": "Delete a journal entry by ID.",
        "input_schema": {
            "type": "object",
            "properties": {
                "je_id": {"type": "integer", "description": "The journal entry ID to delete"},
            },
            "required": ["je_id"],
        },
    },
    {
        "name": "approve_transaction",
        "description": "Approve a transaction, moving it out of the review queue.",
        "input_schema": {
            "type": "object",
            "properties": {
                "transaction_id": {"type": "integer"},
            },
            "required": ["transaction_id"],
        },
    },
    {
        "name": "reject_transaction",
        "description": "Reject a transaction.",
        "input_schema": {
            "type": "object",
            "properties": {
                "transaction_id": {"type": "integer"},
            },
            "required": ["transaction_id"],
        },
    },
    {
        "name": "create_rule",
        "description": (
            "Create a rule in the rules engine that will automatically code or reject future transactions. "
            "Use rule_action='reject' to exclude transactions from the review queue and QBO export (e.g. payroll already handled elsewhere). "
            "Use rule_action='expense' for standard DR/CR coding. "
            "After creating, always call apply_rule with the returned rule_id to apply it to existing pending transactions."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "match_type": {
                    "type": "string",
                    "enum": ["description_contains", "description_exact", "counterparty_contains", "counterparty_exact",
                             "category_equals", "has_category", "kind", "amount_gt", "amount_lt"],
                    "description": "How to match the transaction",
                },
                "match_value": {
                    "type": "string",
                    "description": "The value to match against. Use '*' for has_category.",
                },
                "debit_account": {
                    "type": "string",
                    "description": "Debit account name. Use $source_account or $category as placeholders. Not required for reject action.",
                },
                "credit_account": {
                    "type": "string",
                    "description": "Credit account name. Use $source_account or $category as placeholders. Not required for reject action.",
                },
                "rule_action": {
                    "type": "string",
                    "enum": ["expense", "prepaid", "fixed_asset", "reject"],
                    "description": "What to do when the rule matches. 'reject' excludes the transaction entirely.",
                },
            },
            "required": ["match_type", "match_value", "rule_action"],
        },
    },
    {
        "name": "list_rules",
        "description": "List all existing rules for this client so you can avoid creating duplicates and can reference rule IDs.",
        "input_schema": {
            "type": "object",
            "properties": {},
            "required": [],
        },
    },
    {
        "name": "disable_rule",
        "description": "Disable (deactivate) an existing rule by its ID. Use list_rules first to find the right ID.",
        "input_schema": {
            "type": "object",
            "properties": {
                "rule_id": {"type": "integer"},
            },
            "required": ["rule_id"],
        },
    },
    {
        "name": "apply_rule",
        "description": "Apply an existing rule immediately to all matching pending transactions. Always call this after create_rule.",
        "input_schema": {
            "type": "object",
            "properties": {
                "rule_id": {"type": "integer"},
            },
            "required": ["rule_id"],
        },
    },
    {
        "name": "update_chart_of_accounts",
        "description": (
            "Add one or more account names to the client's Chart of Accounts file. "
            "Use this when the user asks you to add accounts to the COA, or when a transaction "
            "requires an account that doesn't exist yet. The accounts will be merged with any "
            "existing ones and saved sorted alphabetically."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "accounts": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "List of account names to add to the Chart of Accounts.",
                },
            },
            "required": ["accounts"],
        },
    },
]


def _execute_tool(tool_name: str, tool_input: dict, db: Session, client_id: int) -> str:
    """Execute a tool call and return a result string."""
    try:
        if tool_name == "create_journal_entry":
            txn_id = tool_input["transaction_id"]
            txn = db.query(models.Transaction).filter(
                models.Transaction.id == txn_id,
                models.Transaction.client_id == client_id,
            ).first()
            if not txn:
                return f"Transaction {txn_id} not found."
            je_date = None
            if tool_input.get("je_date"):
                from datetime import datetime as _dt
                je_date = _dt.strptime(tool_input["je_date"], "%Y-%m-%d")
            je = models.JournalEntry(
                transaction_id=txn_id,
                debit_account=tool_input["debit_account"],
                credit_account=tool_input["credit_account"],
                amount=tool_input["amount"],
                memo=tool_input.get("memo"),
                je_date=je_date,
                ai_reasoning="Created by AI assistant",
            )
            db.add(je)
            db.flush()
            db.refresh(je)
            db.commit()
            return f"Journal entry created with ID {je.id}: DR {je.debit_account} / CR {je.credit_account} ${je.amount:,.2f}."

        elif tool_name == "update_journal_entry":
            je_id = tool_input["je_id"]
            je = (
                db.query(models.JournalEntry)
                .join(models.Transaction)
                .filter(models.JournalEntry.id == je_id, models.Transaction.client_id == client_id)
                .first()
            )
            if not je:
                return f"Journal entry {je_id} not found."
            if "debit_account" in tool_input:
                je.debit_account = tool_input["debit_account"]
            if "credit_account" in tool_input:
                je.credit_account = tool_input["credit_account"]
            if "amount" in tool_input:
                je.amount = tool_input["amount"]
            if "memo" in tool_input:
                je.memo = tool_input["memo"]
            if "je_date" in tool_input:
                from datetime import datetime as _dt
                je.je_date = _dt.strptime(tool_input["je_date"], "%Y-%m-%d")
            db.commit()
            return f"Journal entry {je_id} updated successfully."

        elif tool_name == "delete_journal_entry":
            je_id = tool_input["je_id"]
            je = (
                db.query(models.JournalEntry)
                .join(models.Transaction)
                .filter(models.JournalEntry.id == je_id, models.Transaction.client_id == client_id)
                .first()
            )
            if not je:
                return f"Journal entry {je_id} not found."
            db.delete(je)
            db.commit()
            return f"Journal entry {je_id} deleted."

        elif tool_name == "approve_transaction":
            txn_id = tool_input["transaction_id"]
            txn = db.query(models.Transaction).filter(
                models.Transaction.id == txn_id,
                models.Transaction.client_id == client_id,
            ).first()
            if not txn:
                return f"Transaction {txn_id} not found."
            txn.status = models.TransactionStatus.approved
            db.commit()
            return f"Transaction {txn_id} approved."

        elif tool_name == "reject_transaction":
            txn_id = tool_input["transaction_id"]
            txn = db.query(models.Transaction).filter(
                models.Transaction.id == txn_id,
                models.Transaction.client_id == client_id,
            ).first()
            if not txn:
                return f"Transaction {txn_id} not found."
            txn.status = models.TransactionStatus.rejected
            db.commit()
            return f"Transaction {txn_id} rejected."

        elif tool_name == "create_rule":
            import rules_engine as _re
            rule_action = tool_input.get("rule_action", "expense")
            debit = tool_input.get("debit_account", "Uncoded")
            credit = tool_input.get("credit_account", "Uncoded")
            rule = models.Rule(
                client_id=client_id,
                match_type=tool_input["match_type"],
                match_value=tool_input["match_value"],
                debit_account=debit,
                credit_account=credit,
                rule_action=rule_action,
                active=True,
            )
            db.add(rule)
            db.flush()
            db.refresh(rule)
            db.commit()
            action_desc = "REJECT (excluded from QBO)" if rule_action == "reject" else f"DR {debit} / CR {credit}"
            return (
                f"Rule created (id={rule.id}): when {rule.match_type}='{rule.match_value}' → {action_desc}. "
                f"Now call apply_rule with rule_id={rule.id} to apply it to pending transactions."
            )

        elif tool_name == "list_rules":
            rules = db.query(models.Rule).filter(models.Rule.client_id == client_id).all()
            if not rules:
                return "No rules configured yet."
            lines = []
            for r in rules:
                status = "active" if r.active else "disabled"
                action = r.rule_action or "expense"
                if action == "reject":
                    acct = "REJECT"
                else:
                    acct = f"DR {r.debit_account} / CR {r.credit_account}"
                lines.append(f"id={r.id} [{status}] {r.match_type}='{r.match_value}' → {action}: {acct}")
            return "Current rules:\n" + "\n".join(lines)

        elif tool_name == "disable_rule":
            rule_id = tool_input["rule_id"]
            rule = db.query(models.Rule).filter(
                models.Rule.id == rule_id, models.Rule.client_id == client_id
            ).first()
            if not rule:
                return f"Rule {rule_id} not found."
            rule.active = False
            db.commit()
            return f"Rule {rule_id} disabled."

        elif tool_name == "apply_rule":
            import rules_engine as _re
            rule_id = tool_input["rule_id"]
            rule = db.query(models.Rule).filter(
                models.Rule.id == rule_id, models.Rule.client_id == client_id
            ).first()
            if not rule:
                return f"Rule {rule_id} not found."

            pending = db.query(models.Transaction).filter(
                models.Transaction.client_id == client_id,
                models.Transaction.status == models.TransactionStatus.pending,
            ).all()

            applied = 0
            for txn in pending:
                if not _re.match_rule(txn, [rule]):
                    continue
                if rule.rule_action == "reject":
                    txn.status = models.TransactionStatus.rejected
                    for je in list(txn.journal_entries):
                        db.delete(je)
                    applied += 1
                else:
                    for je in list(txn.journal_entries):
                        db.delete(je)
                    for jd in _re.apply_rule_jes(rule, txn):
                        db.add(models.JournalEntry(
                            je_number=models.next_je_number(db),
                            transaction_id=txn.id,
                            debit_account=jd["debit_account"],
                            credit_account=jd["credit_account"],
                            amount=abs(jd.get("amount", txn.amount)),
                            je_date=jd.get("je_date"),
                            memo=jd.get("memo"),
                            rule_applied=rule.id,
                            ai_confidence=jd.get("ai_confidence", 1.0),
                            ai_reasoning=jd.get("ai_reasoning"),
                            is_recurring=jd.get("is_recurring", False),
                        ))
                    applied += 1
            db.commit()
            verb = "rejected" if rule.rule_action == "reject" else "coded"
            return f"Rule {rule_id} applied: {verb} {applied} pending transaction(s)."

        elif tool_name == "update_chart_of_accounts":
            new_accounts = [a.strip() for a in tool_input.get("accounts", []) if a.strip()]
            if not new_accounts:
                return "No account names provided."

            client = db.query(models.Client).filter(models.Client.id == client_id).first()
            if not client:
                return "Client not found."

            existing: list[str] = []
            ref = client.chart_of_accounts_path
            if ref:
                existing_text = storage.read_text(ref)
                if existing_text is not None:
                    if not ref.endswith(".txt") and not ref.split("?")[0].endswith(".txt"):
                        return "COA file is not .txt format — can only append to .txt files. Ask the user to re-upload as a .txt file first."
                    existing = [l.strip() for l in existing_text.splitlines() if l.strip()]

            merged = sorted(set(existing) | set(new_accounts), key=str.casefold)
            contents = ("\n".join(merged) + "\n").encode()
            filename = f"coa_client_{client_id}.txt"
            new_ref = storage.upload(filename, contents)

            if not ref:
                client.chart_of_accounts_path = new_ref
                db.commit()

            added = set(new_accounts) - set(existing)
            return f"Chart of Accounts updated. Added {len(added)} account(s): {', '.join(sorted(added))}. Total accounts: {len(merged)}."

        return f"Unknown tool: {tool_name}"
    except Exception as e:
        return f"Tool error: {e}"


# ── Endpoint ──────────────────────────────────────────────────────────────────

_HISTORY_LIMIT = 50


@router.post("/", response_model=ChatResponse)
def chat(
    payload: ChatRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="ANTHROPIC_API_KEY is not configured on the server.",
        )

    client = (
        db.query(models.Client)
        .filter(models.Client.id == payload.client_id, models.Client.user_id == current_user.id)
        .first()
    )
    if not client:
        raise HTTPException(status_code=404, detail="Client not found.")

    system_prompt = _build_system_prompt(client, db, payload.current_page, payload.page_context)

    # Load persisted history
    db_history = (
        db.query(models.ClientChatMessage)
        .filter(models.ClientChatMessage.client_id == client.id)
        .order_by(models.ClientChatMessage.created_at.asc())
        .all()
    )
    history_messages = [{"role": m.role, "content": m.content} for m in db_history[-_HISTORY_LIMIT:]]
    session_messages = [{"role": m.role, "content": m.content} for m in payload.messages]

    # Attach images to the last user message if provided
    if payload.images and session_messages and session_messages[-1]["role"] == "user":
        text = session_messages[-1]["content"]
        content: list = [{"type": "text", "text": text}] if text else []
        for img in payload.images:
            content.append({
                "type": "image",
                "source": {"type": "base64", "media_type": img.media_type, "data": img.data},
            })
        session_messages[-1] = {"role": "user", "content": content}

    messages = history_messages + session_messages

    anthropic_client = anthropic.Anthropic(api_key=api_key)

    # Agentic loop: keep going until Claude stops using tools
    tool_results_summary: list[str] = []
    while True:
        response = anthropic_client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=2048,
            system=system_prompt,
            tools=_TOOLS,
            messages=messages,
        )

        if response.stop_reason == "tool_use":
            # Serialize content blocks to plain dicts for the next API call
            assistant_content = []
            for block in response.content:
                if block.type == "tool_use":
                    assistant_content.append({
                        "type": "tool_use",
                        "id": block.id,
                        "name": block.name,
                        "input": block.input,
                    })
                elif block.type == "text":
                    assistant_content.append({"type": "text", "text": block.text})

            tool_results = []
            for block in response.content:
                if block.type == "tool_use":
                    result = _execute_tool(block.name, block.input, db, client.id)
                    tool_results_summary.append(f"{block.name}: {result}")
                    tool_results.append({
                        "type": "tool_result",
                        "tool_use_id": block.id,
                        "content": result,
                    })

            messages.append({"role": "assistant", "content": assistant_content})
            messages.append({"role": "user", "content": tool_results})
        else:
            # Final text response
            reply = next((b.text for b in response.content if hasattr(b, "text")), "Done.")
            break

    # Persist the user message and assistant reply.
    # If the user sent images, generate a text description to store instead of raw base64.
    if session_messages:
        user_content = session_messages[-1]["content"]
        if payload.images and isinstance(user_content, list):
            try:
                desc_content: list = []
                for img in payload.images:
                    desc_content.append({
                        "type": "image",
                        "source": {"type": "base64", "media_type": img.media_type, "data": img.data},
                    })
                desc_content.append({
                    "type": "text",
                    "text": "Describe what is shown in this image in one concise sentence, for use as a conversation history summary.",
                })
                desc_resp = anthropic_client.messages.create(
                    model="claude-haiku-4-5-20251001",
                    max_tokens=100,
                    messages=[{"role": "user", "content": desc_content}],
                )
                image_desc = next((b.text for b in desc_resp.content if hasattr(b, "text")), "image attachment").strip()
            except Exception:
                image_desc = "image attachment"
            text_part = next((b["text"] for b in user_content if isinstance(b, dict) and b.get("type") == "text"), "").strip()
            stored_content = f"{text_part}\n[Image: {image_desc}]".strip() if text_part else f"[Image: {image_desc}]"
        else:
            stored_content = user_content if isinstance(user_content, str) else str(user_content)
        db.add(models.ClientChatMessage(
            client_id=client.id,
            role=session_messages[-1]["role"],
            content=stored_content,
        ))
    db.add(models.ClientChatMessage(
        client_id=client.id,
        role="assistant",
        content=reply,
    ))
    db.commit()

    return ChatResponse(reply=reply)


@router.get("/history")
def get_chat_history(
    client_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Return persisted chat history for a client so the frontend can restore it on load."""
    client = (
        db.query(models.Client)
        .filter(models.Client.id == client_id, models.Client.user_id == current_user.id)
        .first()
    )
    if not client:
        raise HTTPException(status_code=404, detail="Client not found.")

    messages = (
        db.query(models.ClientChatMessage)
        .filter(models.ClientChatMessage.client_id == client_id)
        .order_by(models.ClientChatMessage.created_at.asc())
        .all()
    )
    return [{"role": m.role, "content": m.content} for m in messages[-_HISTORY_LIMIT:]]
